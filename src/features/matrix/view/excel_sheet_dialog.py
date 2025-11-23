# src/features/matrix/view/excel_sheet_dialog.py
from PyQt5.QtWidgets import (QDialog, QVBoxLayout, QHBoxLayout, QLabel, 
                             QComboBox, QPushButton, QMessageBox)
from PyQt5.QtCore import Qt

from src.features.matrix.service.document_parsers.excel_parser import ExcelParser


class ExcelSheetDialog(QDialog):
    """Excel工作表选择对话框 - 用于导入Excel文件时选择工作表"""
    
    def __init__(self, file_path, parent=None):
        super().__init__(parent)
        self.file_path = file_path
        self.sheet_names = []
        self.selected_sheet = None
        
        self.setWindowTitle("选择工作表")
        self.setModal(True)
        self.resize(300, 150)
        
        # 获取工作表名称列表
        self._get_sheet_names()
        
        self._setup_ui()
        
    def _get_sheet_names(self):
        """获取Excel文件中的工作表名称列表"""
        try:
            from openpyxl import load_workbook
            
            # 只加载工作簿的工作表名称，不加载数据
            workbook = load_workbook(self.file_path, read_only=True, data_only=True)
            self.sheet_names = workbook.sheetnames
            workbook.close()
        except Exception as e:
            # 如果无法获取工作表名称，则使用默认值
            self.sheet_names = ["Sheet1"]
            
    def _setup_ui(self):
        """设置用户界面"""
        layout = QVBoxLayout()
        
        # 工作表选择
        sheet_layout = QHBoxLayout()
        sheet_layout.addWidget(QLabel("工作表:"))
        self.sheet_combo = QComboBox()
        
        # 添加工作表名称到下拉框
        for sheet_name in self.sheet_names:
            self.sheet_combo.addItem(sheet_name)
            
        # 默认选择第一个工作表
        if self.sheet_names:
            self.selected_sheet = self.sheet_names[0]
            self.sheet_combo.setCurrentText(self.selected_sheet)
            
        # 连接信号
        self.sheet_combo.currentTextChanged.connect(self._on_sheet_changed)
        
        sheet_layout.addWidget(self.sheet_combo)
        layout.addLayout(sheet_layout)
        
        # 按钮
        button_layout = QHBoxLayout()
        self.ok_button = QPushButton("确定")
        self.cancel_button = QPushButton("取消")
        
        self.ok_button.clicked.connect(self._on_ok_clicked)
        self.cancel_button.clicked.connect(self.reject)
        
        button_layout.addWidget(self.ok_button)
        button_layout.addWidget(self.cancel_button)
        layout.addLayout(button_layout)
        
        self.setLayout(layout)
        
    def _on_sheet_changed(self, sheet_name):
        """处理工作表选择变化"""
        self.selected_sheet = sheet_name
        
    def _on_ok_clicked(self):
        """处理确定按钮点击事件"""
        self.accept()
        
    def get_selected_sheet(self):
        """获取选择的工作表名称"""
        return self.selected_sheet