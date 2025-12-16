from openpyxl.utils import get_column_letter
from src.core.logger import logger


class LLCRCRFormulaService:
    """LLCR/CR公式计算服务 - 专门处理Excel公式计算"""
    
    def __init__(self):
        pass
        
    def insert_calculation_formulas(self, ws, current_row, point_array, sample_count,
                                   calculate_start_col, bulk_avg_cell, current_cr_cell, test_type):
        """插入计算公式"""
        for j in range(len(point_array)):
            for i in range(1, sample_count + 1):
                col_idx = calculate_start_col + i - 1
                record_col_idx = i + 3  # 原始记录列索引

                if test_type == "CR" and current_cr_cell:
                    formula = f'=({get_column_letter(record_col_idx)}{current_row + j} - ${bulk_avg_cell})/${current_cr_cell}'
                else:
                    formula = f'={get_column_letter(record_col_idx)}{current_row + j} - ${bulk_avg_cell}'

                ws.cell(row=current_row + j, column=col_idx).value = formula

    def handle_delta_r(self, ws, current_row, point_array, sample_count, delta_r_start_col, 
                       calculate_start_col, step_description, test_type, initial_test_rows, 
                       total_row_offset):
        """处理Delta R计算"""
        # 检查是否为初始步骤
        is_initial = "initial" in step_description.lower() and test_type.lower() in step_description.lower()

        # 获取组名
        group_cell_value = ws.cell(row=10 + total_row_offset, column=1).value
        group_name = group_cell_value.replace("Group ", "") if group_cell_value else "Unknown"

        if is_initial:
            # 初始步骤，Delta R的值等于计算值
            for j in range(len(point_array)):
                for i in range(1, sample_count + 1):
                    calc_col_idx = calculate_start_col + i - 1
                    delta_r_col_idx = delta_r_start_col + i - 1
                    formula = f'={get_column_letter(calc_col_idx)}{current_row + j}'
                    ws.cell(row=current_row + j, column=delta_r_col_idx).value = formula

            # 记录初始测试行
            group_cell_value = ws.cell(row=10 + total_row_offset, column=1).value
            group_name = group_cell_value.replace("Group ", "") if group_cell_value else "Unknown"
            initial_test_rows[group_name] = current_row

        else:
            # 非初始步骤，需要减去初始值
            group_cell_value = ws.cell(row=10 + total_row_offset, column=1).value
            group_name = group_cell_value.replace("Group ", "") if group_cell_value else "Unknown"
            initial_row = initial_test_rows.get(group_name)
            if not initial_row:
                logger.warning(f"组 {group_name} 未找到Initial {test_type}步骤！")
                return

            for j in range(len(point_array)):
                for i in range(1, sample_count + 1):
                    calc_col_idx = calculate_start_col + i - 1
                    delta_r_col_idx = delta_r_start_col + i - 1
                    initial_col_idx = calculate_start_col + i - 1

                    formula = f'={get_column_letter(calc_col_idx)}{current_row + j} - {get_column_letter(initial_col_idx)}{initial_row + j}'
                    ws.cell(row=current_row + j, column=delta_r_col_idx).value = formula
                    
    def insert_statistics_formulas(self, ws, current_row, point_array, sample_count,
                                    calculate_start_col, calculate_end_col, stat_start_col,
                                    delta_r_start_col, is_delta_r_checked, test_type):
        """插入统计公式"""
        if is_delta_r_checked and test_type == "LLCR" and delta_r_start_col:
            data_range = f"{get_column_letter(delta_r_start_col)}{current_row}:{get_column_letter(delta_r_start_col + sample_count - 1)}{current_row + len(point_array) - 1}"
        else:
            data_range = f"{get_column_letter(calculate_start_col)}{current_row}:{get_column_letter(calculate_end_col)}{current_row + len(point_array) - 1}"

        # 插入统计公式
        ws.cell(row=current_row, column=stat_start_col).value = f'=MIN({data_range})'  # Min
        ws.cell(row=current_row, column=stat_start_col + 1).value = f'=MAX({data_range})'  # Max
        ws.cell(row=current_row, column=stat_start_col + 2).value = f'=AVERAGE({data_range})'  # Avg
        ws.cell(row=current_row, column=stat_start_col + 3).value = f'=STDEV({data_range})'  # Stdev