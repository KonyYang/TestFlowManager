from src.features.matrix.service.export.service.base_export_service import BaseExportService
from src.features.matrix.service.export.service.excel_formatting_service import ExcelFormattingService
from openpyxl import Workbook
from src.core.logger import logger


class MatingUnmatingExportService(BaseExportService):
    """Mating/Unmating导出服务"""
    
    def __init__(self, data_model):
        super().__init__(data_model)
        self.formatting_service = ExcelFormattingService()
    
    def export_to_excel(self, file_path):
        """导出Mating/Unmating到Excel"""
        logger.debug(f"开始导出Mating/Unmating到 {file_path}")
        try:
            # 创建工作簿
            wb = Workbook()
            ws = wb.active
            
            # TODO: 实现Mating/Unmating特定的导出逻辑
            # 这里应该添加Mating/Unmating特有的数据处理和格式化
            
            # 示例基本结构
            headers = self.data_model.get_headers()
            rows = self.data_model.get_rows()
            
            # 写入表头
            for col_idx, header in enumerate(headers, 1):
                ws.cell(row=1, column=col_idx, value=header)
            
            # 写入数据行
            for row_idx, row_data in enumerate(rows, 2):  # 从第2行开始
                for col_idx, cell_value in enumerate(row_data):
                    ws.cell(row=row_idx, column=col_idx + 1, value=cell_value)
            
            # 应用格式化
            self.formatting_service.format_worksheet(ws)
            
            # 保存文件
            return self._save_workbook_safely(wb, file_path)
        except Exception as e:
            logger.error(f"导出Mating/Unmating失败: {e}", exc_info=True)
            return False