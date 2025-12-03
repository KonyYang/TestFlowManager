from src.features.matrix.model.matrix_data import MatrixData
from src.core.logger import logger
from src.features.matrix.service.export.service.base_export_service import BaseExportService
# 导入Excel格式化服务
from src.features.matrix.service.export.service.excel_formatting_service import ExcelFormattingService
from openpyxl import Workbook


class MatrixEditorExcelExportService(BaseExportService):
    """Matrix编辑器Excel导出服务 - 处理Matrix编辑器内容导出到Excel的功能"""

    def __init__(self, data_model: MatrixData):
        super().__init__(data_model)
        self.merged_cells_info = []
        # 创建格式化服务实例
        self.formatting_service = ExcelFormattingService()

    def export_to_excel(self, file_path):
        """将Matrix编辑器内容导出到Excel - Service层持久化功能"""
        try:
            # 创建工作簿
            wb = Workbook()
            ws = wb.active

            # 先添加数据行（不包括表头）
            for row_idx, row_data in enumerate(self.data_model.get_rows()):
                for col_idx, cell_value in enumerate(row_data):
                    ws.cell(row=row_idx + 1, column=col_idx + 1, value=cell_value)
            
            # 应用合并单元格
            # 移除合并单元格导出的详细日志
            for merge_info in self.data_model.merged_cells_info:
                top_row = merge_info['top_row'] + 1  # +1 because of 1-based indexing
                left_col = merge_info['left_col'] + 1  # +1 because of 1-based indexing
                bottom_row = top_row + merge_info['row_count'] - 1
                right_col = left_col + merge_info['col_count'] - 1
                
                logger.debug(f"处理合并单元格: top_row={top_row}, left_col={left_col}, "
                           f"bottom_row={bottom_row}, right_col={right_col}")
                
                # 先保存合并区域左上角单元格的值
                top_left_value = ws.cell(row=top_row, column=left_col).value
                
                # 清空整个合并区域的值
                for row in range(top_row, bottom_row + 1):
                    for col in range(left_col, right_col + 1):
                        ws.cell(row=row, column=col, value=None)
                
                # 将原值设置回合并区域的左上角单元格
                ws.cell(row=top_row, column=left_col, value=top_left_value)
                
                # 合并单元格
                ws.merge_cells(
                    start_row=top_row, 
                    start_column=left_col, 
                    end_row=bottom_row, 
                    end_column=right_col
                )

            # 应用Matrix个性化格式化
            # 定义Matrix表格的列宽设置
            matrix_column_widths = {
                1: 20,  # 第1列（Test Item）要宽些
                2: 8,  # 第2列（PARA）要窄些
                4: 20,  # 第4列（Condition）要宽些
                5: 20   # 第5列（Requirement）要宽些
            }
            
            # 使用自定义列宽格式化工作表
            self.formatting_service.format_worksheet_with_custom_widths(ws, matrix_column_widths)
            
            # 为首行和首列应用灰色背景
            self.formatting_service.apply_background_fill(ws, rows=[1], cols=[1])

            # 保存文件
            return self._save_workbook_safely(wb, file_path)
        except Exception as e:
            logger.error(f"导出Excel失败: {e}")
            return False