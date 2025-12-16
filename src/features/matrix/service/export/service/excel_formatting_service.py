from openpyxl.styles import Alignment, Border, Side, PatternFill, Font
from openpyxl.utils import get_column_letter
from src.core.logger import logger


class ExcelFormattingService:
    """
    Excel格式化服务
    提供通用的Excel格式化功能，如边框、对齐、自动换行等
    """

    def __init__(self):
        # 定义常用边框样式
        self.thin_border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )
        
        # 定义常用对齐样式
        self.center_alignment = Alignment(
            horizontal='center',
            vertical='center'
        )
        
        self.wrap_alignment = Alignment(
            horizontal='center',
            vertical='center',
            wrap_text=True
        )
        
        # 定义灰色背景填充
        self.gray_fill = PatternFill(start_color="CCCCCC", end_color="CCCCCC", fill_type="solid")

    def apply_borders(self, worksheet, start_row=1, start_col=1, end_row=None, end_col=None):
        """
        为指定范围的单元格添加边框
        
        Args:
            worksheet: 工作表对象
            start_row: 起始行（从1开始）
            start_col: 起始列（从1开始）
            end_row: 结束行（包含）
            end_col: 结束列（包含）
        """
        try:
            if end_row is None:
                end_row = worksheet.max_row
            if end_col is None:
                end_col = worksheet.max_column
                
            for row in range(start_row, end_row + 1):
                for col in range(start_col, end_col + 1):
                    cell = worksheet.cell(row=row, column=col)
                    cell.border = self.thin_border
                    
            # 移除添加边框的详细日志
        except Exception as e:
            logger.error(f"添加边框时出错: {e}")

    def auto_adjust_column_width(self, worksheet, min_width=10, max_width=50):
        """
        自动调整列宽
        
        Args:
            worksheet: 工作表对象
            min_width: 最小列宽
            max_width: 最大列宽
        """
        try:
            for column in worksheet.columns:
                max_length = 0
                column_letter = get_column_letter(column[0].column)
                
                for cell in column:
                    try:
                        if len(str(cell.value)) > max_length:
                            max_length = len(str(cell.value))
                    except:
                        pass
                        
                adjusted_width = min(max_width, max(min_width, max_length + 2))
                worksheet.column_dimensions[column_letter].width = adjusted_width
                
            logger.debug("已完成列宽自动调整")
        except Exception as e:
            logger.error(f"自动调整列宽时出错: {e}")

    def auto_adjust_column_width_custom(self, worksheet, column_widths):
        """
        根据自定义设置调整列宽
        
        Args:
            worksheet: 工作表对象
            column_widths: 列宽设置字典，键为列索引（从1开始），值为宽度
        """
        try:
            max_col = worksheet.max_column
            
            # 设置各列宽度
            for col in range(1, max_col + 1):
                column_letter = get_column_letter(col)
                if col in column_widths:
                    worksheet.column_dimensions[column_letter].width = column_widths[col]
                    
            # 移除自定义列宽调整的详细日志
        except Exception as e:
            logger.error(f"自定义列宽调整时出错: {e}")

    def auto_adjust_row_height(self, worksheet, base_height=15):
        """
        自动调整行高（基于内容换行）
        
        Args:
            worksheet: 工作表对象
            base_height: 基础行高
        """
        try:
            for row in worksheet.iter_rows():
                max_lines = 1
                for cell in row:
                    if cell.value and isinstance(cell.value, str):
                        lines = str(cell.value).count('\n') + 1
                        max_lines = max(max_lines, lines)
                        
                # 设置行高，基于行中最多行数的单元格
                worksheet.row_dimensions[row[0].row].height = base_height * max_lines
                
            # 移除行高自动调整的详细日志
        except Exception as e:
            logger.error(f"自动调整行高时出错: {e}")

    def apply_text_wrap_and_center(self, worksheet, start_row=1, start_col=1, end_row=None, end_col=None):
        """
        为指定范围的单元格应用文本自动换行和居中对齐
        
        Args:
            worksheet: 工作表对象
            start_row: 起始行（从1开始）
            start_col: 起始列（从1开始）
            end_row: 结束行（包含）
            end_col: 结束列（包含）
        """
        try:
            if end_row is None:
                end_row = worksheet.max_row
            if end_col is None:
                end_col = worksheet.max_column
                
            for row in range(start_row, end_row + 1):
                for col in range(start_col, end_col + 1):
                    cell = worksheet.cell(row=row, column=col)
                    cell.alignment = self.wrap_alignment
                    
            # 移除文本换行和居中对齐的详细日志
        except Exception as e:
            logger.error(f"应用文本换行和居中对齐时出错: {e}")

    def apply_background_fill(self, worksheet, rows=None, cols=None, fill=None):
        """
        为指定行或列应用背景填充
        
        Args:
            worksheet: 工作表对象
            rows: 行索引列表（从1开始）
            cols: 列索引列表（从1开始）
            fill: 填充样式，默认使用灰色填充
        """
        try:
            max_row = worksheet.max_row
            max_col = worksheet.max_column
            fill_style = fill if fill else self.gray_fill
            
            # 为指定行应用背景填充
            if rows:
                for row in rows:
                    if 1 <= row <= max_row:
                        for col in range(1, max_col + 1):
                            cell = worksheet.cell(row=row, column=col)
                            cell.fill = fill_style
                
            # 为指定列应用背景填充
            if cols:
                for col in cols:
                    if 1 <= col <= max_col:
                        for row in range(1, max_row + 1):
                            cell = worksheet.cell(row=row, column=col)
                            cell.fill = fill_style
                
            # 移除背景填充的详细日志
        except Exception as e:
            logger.error(f"应用背景填充时出错: {e}")

    def format_range(self, worksheet, start_row=1, start_col=1, end_row=None, end_col=None, 
                     border=True, wrap_text=True, center=True):
        """
        为指定范围应用综合格式化
        
        Args:
            worksheet: 工作表对象
            start_row: 起始行（从1开始）
            start_col: 起始列（从1开始）
            end_row: 结束行（包含）
            end_col: 结束列（包含）
            border: 是否添加边框
            wrap_text: 是否启用文本换行
            center: 是否居中对齐
        """
        try:
            if end_row is None:
                end_row = worksheet.max_row
            if end_col is None:
                end_col = worksheet.max_column
                
            alignment = None
            if center and wrap_text:
                alignment = self.wrap_alignment
            elif center:
                alignment = self.center_alignment
                
            for row in range(start_row, end_row + 1):
                for col in range(start_col, end_col + 1):
                    cell = worksheet.cell(row=row, column=col)
                    if border:
                        cell.border = self.thin_border
                    if alignment:
                        cell.alignment = alignment
                        
            logger.debug(f"已完成范围 {start_row}:{end_row}, {start_col}:{end_col} 的综合格式化")
        except Exception as e:
            logger.error(f"综合格式化时出错: {e}")

    def format_worksheet(self, worksheet):
        """
        为整个工作表应用标准格式化（通用功能）
        包含：边框、自动换行和居中对齐、自适应列宽和行高
        
        Args:
            worksheet: 工作表对象
        """
        try:
            # 应用边框
            self.apply_borders(worksheet)
            
            # 应用文本换行和居中对齐
            self.apply_text_wrap_and_center(worksheet)
            
            # 自动调整列宽
            self.auto_adjust_column_width(worksheet)
            
            # 自动调整行高
            self.auto_adjust_row_height(worksheet)
            
            logger.debug("已完成工作表标准格式化")
        except Exception as e:
            logger.error(f"工作表标准格式化时出错: {e}")

    def format_worksheet_with_custom_widths(self, worksheet, column_widths):
        """
        为工作表应用标准格式化并使用自定义列宽
        
        Args:
            worksheet: 工作表对象
            column_widths: 列宽设置字典，键为列索引（从1开始），值为宽度
        """
        try:
            # 应用边框
            self.apply_borders(worksheet)
            
            # 应用文本换行和居中对齐
            self.apply_text_wrap_and_center(worksheet)
            
            # 应用自定义列宽
            self.auto_adjust_column_width_custom(worksheet, column_widths)
            
            # 自动调整行高
            self.auto_adjust_row_height(worksheet)
            
            # 移除工作表自定义列宽格式化的详细日志
        except Exception as e:
            logger.error(f"工作表自定义列宽格式化时出错: {e}")

    def format_worksheet_with_background(self, worksheet, header_rows=None, header_cols=None):
        """
        为工作表应用标准格式化并添加背景色
        
        Args:
            worksheet: 工作表对象
            header_rows: 需要添加背景色的行索引列表（从1开始）
            header_cols: 需要添加背景色的列索引列表（从1开始）
        """
        try:
            # 应用标准格式化
            self.format_worksheet(worksheet)
            
            # 应用背景色
            self.apply_background_fill(worksheet, header_rows, header_cols)
            
            logger.debug("已完成工作表带背景色格式化")
        except Exception as e:
            logger.error(f"工作表带背景色格式化时出错: {e}")
            
    def format_range_bold_header(self, ws, title_row, stat_start_col):
        """
        设置表头样式（加粗字体和灰色背景）
        
        Args:
            ws: 工作表对象
            title_row: 标题行号
            stat_start_col: 统计列起始列号
        """
        # 设置表头样式
        header_font = Font(name='Arial', size=9, bold=True)
        header_fill = PatternFill(start_color="DCDCDC", end_color="DCDCDC", fill_type="solid")
        
        # 设置表头行的字体和背景色
        for col in range(1, ws.max_column + 1):
            cell = ws.cell(row=title_row, column=col)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
            
    def format_merge_cells_style(self, ws, start_row, start_col, row_span):
        """
        设置合并单元格样式
        
        Args:
            ws: 工作表对象
            start_row: 起始行号
            start_col: 起始列号
            row_span: 合并行数
        """
        # 实现合并单元格的样式设置
        for row in range(start_row, start_row + row_span):
            cell = ws.cell(row=row, column=start_col)
            cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)