from src.features.matrix.model.matrix_data import MatrixData
from src.features.ltr_manager.model.ltr_application_data import LTRApplicationData
from src.core.logger import logger
from src.features.matrix.service.export.service.excel_formatting_service import ExcelFormattingService
from openpyxl import Workbook


class TestStatusExportService:
    """Test Status表导出服务 - 生成特定格式的Test Status表"""

    def __init__(self, data_model, ltr_data: LTRApplicationData = None):
        self.data_model = data_model
        self.ltr_data = ltr_data
        # 创建格式化服务实例
        self.formatting_service = ExcelFormattingService()
        logger.debug(f"TestStatusExportService 初始化完成，ltr_data: {ltr_data}")

    def export_test_status_to_excel(self, file_path):
        """将Matrix数据导出为Test Status表格式"""
        logger.debug(f"开始导出Test Status表到 {file_path}")
        try:
            # 创建工作簿
            wb = Workbook()
            ws = wb.active
            logger.debug("创建工作簿成功")

            # 确定截止行（找到第一列包含"Sample"的行）
            cutoff_row_index = self._find_sample_row_index()
            logger.debug(f"找到Sample行索引: {cutoff_row_index}")
            
            # 获取处理后的数据（移除指定列）
            processed_headers, processed_rows = self._process_matrix_data(cutoff_row_index)
            logger.debug(f"处理后的表头数量: {len(processed_headers)}, 数据行数: {len(processed_rows)}")
            
            # 写入表头
            for col_idx, header in enumerate(processed_headers):
                ws.cell(row=1, column=col_idx + 1, value=header)
            logger.debug("写入表头完成")
            
            # 写入数据行
            for row_idx, row_data in enumerate(processed_rows):
                for col_idx, cell_value in enumerate(row_data):
                    ws.cell(row=row_idx + 2, column=col_idx + 1, value=cell_value)
            logger.debug("写入数据行完成")
            
            # 添加"Estimated completion date in Clarizen"行
            if self.ltr_data and self.ltr_data.estimated_completion_date:
                est_date_row = len(processed_rows) + 2  # +2是因为有表头行和1-based索引
                ws.cell(row=est_date_row, column=1, value="Estimated completion date in Clarizen")
                # 合并该行除第一列外的所有列
                if len(processed_headers) > 1:
                    ws.merge_cells(start_row=est_date_row, start_column=2, 
                                   end_row=est_date_row, end_column=len(processed_headers))
                    ws.cell(row=est_date_row, column=2, value=self.ltr_data.estimated_completion_date)
                logger.debug(f"添加Estimated completion date行，值: {self.ltr_data.estimated_completion_date}")
            
            # 应用格式化
            self._apply_formatting(ws, len(processed_rows) + 2, len(processed_headers))
            logger.debug("应用格式化完成")
            
            # 保存文件
            wb.save(file_path)
            logger.debug(f"文件保存成功: {file_path}")
            return True
        except PermissionError:
            # 文件被其他程序占用（如Excel）
            logger.error(f"导出Test Status表失败: 文件被占用，可能已在Excel中打开")
            return False
        except Exception as e:
            logger.error(f"导出Test Status表失败: {e}", exc_info=True)
            return False

    def _find_sample_row_index(self):
        """找到第一列包含'Sample'的行索引"""
        logger.debug("开始查找Sample行索引")
        rows = self.data_model.get_rows()
        logger.debug(f"总行数: {len(rows)}")
        for row_idx, row_data in enumerate(rows):
            logger.debug(f"检查第{row_idx}行数据: {row_data}")
            if row_data and len(row_data) > 0 and "Sample" in str(row_data[0]):
                logger.debug(f"在第{row_idx}行找到Sample")
                return row_idx
        # 如果没找到，返回所有行
        logger.debug("未找到Sample行，返回所有行")
        return len(rows)

    def _process_matrix_data(self, cutoff_row_index):
        """处理矩阵数据，移除指定列并添加Status列"""
        logger.debug(f"开始处理矩阵数据，截止行索引: {cutoff_row_index}")
        # 原始表头
        original_headers = self.data_model.get_headers()[:]
        logger.debug(f"原始表头: {original_headers}")
        
        # 处理空表头的情况
        if not original_headers:
            logger.debug("表头为空，返回空结果")
            return [], []
        
        # 确定要移除的列索引 (2,3,4,5列对应索引1,2,3,4，以及最后一列Notes)
        columns_to_remove = {1, 2, 3, 4}
        # 只有当有足够的列时才移除最后一列
        if len(original_headers) > 1:
            columns_to_remove.add(len(original_headers) - 1)
        
        logger.debug(f"要移除的列索引: {columns_to_remove}")
        
        # 处理表头
        processed_headers = []
        for idx, header in enumerate(original_headers):
            if idx not in columns_to_remove:
                processed_headers.append(header)
        logger.debug(f"处理后的表头(移除特定列后): {processed_headers}")
        
        # 添加Status列
        processed_headers.append("Status")
        logger.debug(f"添加Status列后的表头: {processed_headers}")
        
        # 处理数据行
        processed_rows = []
        rows = self.data_model.get_rows()
        logger.debug(f"原始数据行数: {len(rows)}")
        for row_idx, row_data in enumerate(rows[:cutoff_row_index]):
            logger.debug(f"处理第{row_idx}行数据: {row_data}")
            processed_row = []
            # 确保row_data不为空
            if row_data:
                for col_idx, cell_value in enumerate(row_data):
                    if col_idx not in columns_to_remove:
                        processed_row.append(cell_value)
            # 添加Status列的默认值
            processed_row.append("No Start")
            processed_rows.append(processed_row)
            logger.debug(f"处理后的第{row_idx}行: {processed_row}")
            
        logger.debug(f"处理完成，最终数据行数: {len(processed_rows)}")
        return processed_headers, processed_rows

    def _apply_formatting(self, worksheet, data_rows, data_cols):
        """应用表格格式化"""
        try:
            logger.debug(f"开始应用格式化，数据行数: {data_rows}, 数据列数: {data_cols}")
            # 定义列宽设置
            column_widths = {
                1: 20,  # 第1列要宽些
                data_cols: 15   # Status列
            }
            
            # 使用自定义列宽格式化工作表
            self.formatting_service.format_worksheet_with_custom_widths(worksheet, column_widths)
            
            # 为首行应用灰色背景
            self.formatting_service.apply_background_fill(worksheet, rows=[1])
            
            # 应用边框到整个数据区域
            self.formatting_service.apply_borders(worksheet, 1, 1, data_rows, data_cols)
            logger.debug("格式化应用完成")
            
        except Exception as e:
            logger.error(f"应用格式化时出错: {e}", exc_info=True)